"""nb_17_expert_sheets - build BLINDED Excel sheets for the human review steps (R1-4 fine-grained, R3 narrator hallucination).

(1) narrator_rater_A.xlsx / narrator_rater_B.xlsx : the SAME stratified sample of narrator texts for two independent raters.
    Raters see evidence + text only (no condition label, no judge output). Key saved separately (narrator_key.csv).
(2) finegrained_expert.xlsx : one row per cluster; the author's FINE-GRAINED labels (top 3 with fractions) and the four methods' calls
    in a random, unlabelled order (call_1..call_4); the expert marks Y/P/N for each call. Key saved separately (finegrained_key.csv).

Runner:
    DATASETS = ["GSE131907", "GSE132465", "GSE176078"]
    N_PER_DATASET = 20          # narrator texts per dataset (half constrained, half unconstrained); 60 in total
    exec(open("/lakehouse/default/Files/scRNA/code/nb_17_expert_sheets.py").read())
Needs: reanalysis/finegrained_template.csv (nb_12), reanalysis/halluc_manual_review_<DS>.csv (nb_13), h5ad/<DS>_clustered.h5ad.
Output folder: Files/scRNA/reanalysis/expert/
"""
import json, random, re
from pathlib import Path
import numpy as np, pandas as pd

DATASETS = globals().get("DATASETS", ["GSE131907", "GSE132465", "GSE176078"])
N_PER_DATASET = globals().get("N_PER_DATASET", 20)
FINE = dict(globals().get("FINE", {}))
FINE_DEFAULT = {"GSE131907": ["Cell_type.refined", "Cell_subtype"], "GSE132465": ["Cell_subtype"], "GSE176078": ["celltype_minor", "celltype_subset"]}
ROOT = Path("/lakehouse/default/Files/scRNA"); REA = ROOT/"reanalysis"; OUT = REA/"expert"; OUT.mkdir(parents=True, exist_ok=True)
METHODS = ["scINTERP", "CellTypist", "naive_LLM", "SingleR"]

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation

def write_book(path, instructions, df, widths, wrap_cols, dv_cols=None, dv_list='"Y,N"'):
    wb = Workbook(); ws0 = wb.active; ws0.title = "INSTRUCTIONS"
    for i, line in enumerate(instructions, 1):
        ws0.cell(i, 1, line).alignment = Alignment(wrap_text=True, vertical="top")
    ws0.column_dimensions["A"].width = 130
    ws = wb.create_sheet("REVIEW"); ws.append(list(df.columns))
    for c in ws[1]: c.font = Font(bold=True)
    for row in df.itertuples(index=False):
        ws.append([("" if (isinstance(v, float) and np.isnan(v)) else v) for v in row])
    for j, name in enumerate(df.columns, 1):
        ws.column_dimensions[ws.cell(1, j).column_letter].width = widths.get(name, 14)
        if name in wrap_cols:
            for r in range(2, len(df)+2): ws.cell(r, j).alignment = Alignment(wrap_text=True, vertical="top")
    if dv_cols:
        dv = DataValidation(type="list", formula1=dv_list, allow_blank=True); ws.add_data_validation(dv)
        for name in dv_cols:
            j = list(df.columns).index(name)+1; L = ws.cell(1, j).column_letter
            dv.add(f"{L}2:{L}{len(df)+1}")
    ws.freeze_panes = "A2"; wb.save(path)

# ============================== (1) narrator sheets ==============================
frames = []
for ds in DATASETS:
    f = REA/f"halluc_manual_review_{ds}.csv"
    if f.exists(): frames.append(pd.read_csv(f))
    else: print("missing", f.name)
M = pd.concat(frames, ignore_index=True)
rng = np.random.default_rng(0); picked = []
for ds in DATASETS:
    for cond in ["constrained", "unconstrained"]:
        sub = M[(M.dataset == ds) & (M.condition == cond)]
        k = min(len(sub), N_PER_DATASET//2)
        picked.append(sub.iloc[rng.permutation(len(sub))[:k]])
S = pd.concat(picked, ignore_index=True)
S = S.iloc[rng.permutation(len(S))].reset_index(drop=True)
S["item_id"] = [f"N{i+1:03d}" for i in range(len(S))]
assert S.evidence.str.len().max() < 32000, "an evidence cell exceeds the Excel cell limit"
key = S[["item_id", "dataset", "condition", "cluster", "run", "judge_unsupported"]]; key.to_csv(OUT/"narrator_key.csv", index=False)
N_INSTR = [
 "TASK - narrator review. For each row read EVIDENCE (JSON given to a writing model) and TEXT (what it wrote).",
 "An UNSUPPORTED claim is a statement of fact in TEXT about a marker gene, pathway, disease association, cell-type call, confidence, or flag that is NOT in EVIDENCE or contradicts it.",
 "Do NOT count: (a) generic biological explanation that adds no new specific fact from outside the evidence, (b) hedging words, (c) the writer suggesting validation markers to check next (record those in column suggested_markers_only instead).",
 "Count each distinct unsupported claim once. Fill: n_unsupported (integer, 0 if none), any_unsupported (Y/N), suggested_markers_only (Y/N: the ONLY extra genes mentioned are 'markers to check').",
 "Work alone and do NOT discuss with the other rater. Do not try to guess which model/condition wrote a text.",
 "Save as the same file name with _DONE added (e.g. narrator_rater_A_DONE.xlsx) and put it in Files/scRNA/reanalysis/expert/.",
]
for tag in ["A", "B"]:
    T = S[["item_id", "dataset", "cluster", "evidence", "text"]].copy()
    if tag == "B": T = T.iloc[np.random.default_rng(1).permutation(len(T))].reset_index(drop=True)   # different order for rater B
    T["n_unsupported"] = ""; T["any_unsupported"] = ""; T["suggested_markers_only"] = ""; T["notes"] = ""
    write_book(OUT/f"narrator_rater_{tag}.xlsx", N_INSTR, T,
               widths={"item_id": 9, "dataset": 11, "cluster": 8, "evidence": 70, "text": 70, "n_unsupported": 14, "any_unsupported": 14, "suggested_markers_only": 18, "notes": 30},
               wrap_cols={"evidence", "text", "notes"}, dv_cols=["any_unsupported", "suggested_markers_only"])
print(f"narrator sheets: {len(S)} texts ({S.groupby(['dataset','condition']).size().to_dict()}) -> narrator_rater_A.xlsx, narrator_rater_B.xlsx")

# ============================== (2) fine-grained sheet ==============================
import anndata as _anndata
tpl = pd.read_csv(REA/"finegrained_template.csv"); tpl["cluster"] = tpl["cluster"].astype(str)
rows, keyrows = [], []
rng2 = random.Random(0)
for ds in DATASETS:
    ad = _anndata.read_h5ad(ROOT/"h5ad"/f"{ds}_clustered.h5ad"); obs = ad.obs
    fcols = [c for c in (FINE.get(ds) or FINE_DEFAULT.get(ds, ["Cell_subtype"])) if c in obs.columns]
    print(ds, "fine-grained author columns used:", fcols)
    for _, r in tpl[tpl.dataset == ds].iterrows():
        cl = r["cluster"]; sub = obs[obs["leiden"].astype(str) == cl]
        d = dict(dataset=ds, cluster=cl, n_cells=len(sub), author_major=r["author"])
        parts = []
        for c in fcols:
            vc = sub[c].astype(str).replace({"nan": "unannotated", "NaN": "unannotated", "None": "unannotated"}).value_counts(normalize=True).head(3)
            parts.append(f"[{c}] " + "; ".join(f"{k} ({v:.2f})" for k, v in vc.items()))
        d["author_fine"] = chr(10).join(parts)
        order = METHODS[:]; rng2.shuffle(order)
        for i, m in enumerate(order, 1):
            d[f"call_{i}"] = r.get(f"call_{m}"); keyrows.append(dict(dataset=ds, cluster=cl, slot=i, method=m))
        rows.append(d)
F = pd.DataFrame(rows)
for i in range(1, 5): F[f"curated_{i}"] = ""
F["notes"] = ""
pd.DataFrame(keyrows).to_csv(OUT/"finegrained_key.csv", index=False)
FG_INSTR = [
 "TASK - fine-grained annotation check. One row per cluster. author_major = the authors' broad label; author_fine = the authors' finer labels (one line per annotation column; 'unannotated' = cells the authors did not subtype) of the cells in that cluster (top 3 with fraction).",
 "call_1 .. call_4 are four different annotation methods' calls for the cluster, in random order (the methods are hidden on purpose).",
 "For each call put Y if it names the SAME cell type / subtype / state as the dominant author fine label (synonyms and accepted equivalents count, e.g. 'CD8+ T cell' = 'CD8 T'), or a clinically/biologically equivalent term at that resolution; put N otherwise.",
 "P = the call names the correct PARENT / broader cell type but not the author's fine subtype or state (e.g. 'Macrophage' where the author label is 'SPP1+ macrophage', or 'T cell' where it is 'Treg'). Use P instead of N in that case. P applies ONLY when the call is a BROADER category that contains the author's label; a different sibling subtype (e.g. NKT where the author label is CD4+ Th) is N.",
 "N = wrong cell type or wrong lineage. We report strict accuracy (Y only) and tolerant accuracy (Y or P).",
 "If the cluster is genuinely mixed (second author label >= 0.4), accept a call that matches either of the top two labels and say so in notes.",
 "Fill curated_1 .. curated_4 (Y, P or N) for every row. Save as finegrained_expert_DONE.xlsx in Files/scRNA/reanalysis/expert/.",
]
wid = {"dataset": 11, "cluster": 8, "n_cells": 8, "author_major": 18, "notes": 40}
for i in range(1, 5): wid[f"call_{i}"] = 30; wid[f"curated_{i}"] = 11
wid["author_fine"] = 70
write_book(OUT/"finegrained_expert.xlsx", FG_INSTR, F, widths=wid, wrap_cols={c for c in F.columns if c.startswith(("author_fine", "call_", "notes"))},
           dv_cols=[f"curated_{i}" for i in range(1, 5)], dv_list='"Y,P,N"')
print(f"fine-grained sheet: {len(F)} clusters -> finegrained_expert.xlsx  (key: finegrained_key.csv - do NOT give the key to the expert)")
print(); print("Files in", OUT)
for p in sorted(OUT.glob("*")): print("  ", p.name)
print(); print("Give raters ONLY the *.xlsx files (not the key files). NOTEBOOK 17 COMPLETE")
